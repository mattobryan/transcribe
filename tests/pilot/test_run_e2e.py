"""End-to-end pilot run with a synthetic aligner and a fake Whisper (no downloads)."""
import json
import sys

import numpy as np
import pytest
from docx import Document

from src.transcribe.pilot import asr as asr_mod
from src.transcribe.pilot import ctc_align, run
from src.transcribe.pilot.textprep import align_form

SR = 16000
LETTERS = "abcdefghijklmnopqrstuvwxyz'"
VOCAB = {"<blank>": 0, **{ch: i + 1 for i, ch in enumerate(LETTERS)}}

TURNS = [
    ("I", "Habari yako, how are you today?", True),
    ("R", "Niko fine sana, nimedownload the form jana.", True),
    ("I", "This whole sentence was summarised by the transcriber and never spoken.", False),
    ("R", "Tulienda kanisa jana and the sermon was about patience.", True),
    ("I", "Sawa, asante sana for your time.", True),
]


class FakeEmitter:
    def __init__(self, *_, **__):
        self.vocab = VOCAB
        self.blank = 0

    @property
    def alphabet(self):
        return list(LETTERS)

    def token_ids(self, text):
        return [VOCAB[c] for c in text if c in VOCAB]


def synth_logp(rng):
    """Speak the verbatim turns (and an untranscribed burst) as CTC-like frames."""
    labels = [0] * 100
    truth = {}
    for n, (_, text, spoken) in enumerate(TURNS):
        if n == 3:                                     # untranscribed speech before turn 3
            labels += [VOCAB[c] for c in "zzqqxxzzqq" for _ in range(3)] + [0] * 60
        if not spoken:
            continue
        for word in text.split():
            letters = align_form(word)
            if not letters:
                continue
            start = len(labels)
            for ch in letters:
                labels += [VOCAB[ch]] * 3 + [0]
            truth[(n, word)] = (start, len(labels) - 1)
            labels += [0] * 8
        labels += [0] * 100
    probs = rng.uniform(0.0, 0.004, size=(len(labels), len(VOCAB)))
    for t, lab in enumerate(labels):
        probs[t, lab] = 0.8
    probs /= probs.sum(axis=1, keepdims=True)
    return np.log(probs).astype(np.float32), truth


class FakeWhisper:
    segments = []

    def __init__(self, model, **_):
        self.name = model

    def transcribe(self, clip, language):
        start = float(clip[0])                          # audio is a time ramp
        seg = min(self.segments, key=lambda s: abs(s["start"] - start))
        if abs(seg["start"] - start) > 0.01:            # a silence gap
            return {"text": "", "detected_language": "sw", "language_probability": 1.0}
        words = seg["ref_words"]
        if language == "en":
            words = [w for i, w in enumerate(words) if i % 2 == 0]
        return {"text": " ".join(words), "detected_language": language or "sw",
                "language_probability": 0.9}


@pytest.fixture
def pilot(tmp_path, monkeypatch):
    rng = np.random.default_rng(0)
    logp, truth = synth_logp(rng)
    doc = Document()
    doc.add_paragraph("Header line")
    for speaker, text, _ in TURNS:
        doc.add_paragraph(f"{speaker}: {text}")
    docx_path = tmp_path / "rec.docx"
    doc.save(docx_path)
    seconds = logp.shape[0] * ctc_align.FRAME_SECONDS
    ramp = (np.arange(int(seconds * SR)) / SR).astype(np.float32)

    monkeypatch.setattr(ctc_align, "Emitter", FakeEmitter)
    monkeypatch.setattr(ctc_align, "load_or_compute_emissions", lambda *a, **k: logp)
    monkeypatch.setattr(run, "load_audio", lambda path: ramp)
    monkeypatch.setattr(asr_mod, "WhisperRunner", FakeWhisper)
    out = tmp_path / "out"
    argv = ["run", "--audio", "rec.wav", "--transcript", str(docx_path), "--out", str(out),
            "--languages", "sw,en", "--export-review"]

    def go():
        monkeypatch.setattr(sys, "argv", argv)
        original = run.build_segments

        def capture(*a, **k):
            segs = original(*a, **k)
            FakeWhisper.segments = segs
            return segs
        monkeypatch.setattr(run, "build_segments", capture)
        run.main()
        return out
    return go, truth


def test_pipeline(pilot):
    go, truth = pilot
    out = go()
    alignment = json.loads((out / "alignment.json").read_text())
    turns = alignment["turns"]
    assert turns[2]["confidence"] < 0.5, "non-verbatim turn must be flagged"
    for n in (0, 1, 3, 4):
        assert turns[n]["confidence"] > 0.6, f"turn {n} should recover after the bad turn"
    # word boundaries within one frame of the truth for spoken turns
    words = alignment["words"]
    for w in words:
        key = (w["turn_index"], w["raw"])
        if key in truth and w["score"] is not None and w["turn_index"] != 2:
            start, end = truth[key]
            assert abs(w["start"] / 0.02 - start) <= 1 and abs(w["end"] / 0.02 - end) <= 1, key

    report = json.loads((out / "report.json").read_text())
    by_lang = {e["language"]: e for e in report["asr"]}
    assert by_lang["sw"]["clean"]["wer"] == 0.0
    assert by_lang["en"]["clean"]["wer"] > 0.3
    assert report["alignment"]["turns_below_0_5"] == 1
    assert (out / "report.md").read_text().startswith("# Pilot report")

    review = json.loads((out / "review_manifest.json").read_text())
    assert review["segments"] and all(s["transcript"] for s in review["segments"])
    assert all((out / "chunks" / f"{s['segment_id']}.wav").exists() for s in review["segments"])

    # second run reuses every cache
    go()
