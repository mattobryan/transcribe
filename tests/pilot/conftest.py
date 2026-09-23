import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


@pytest.fixture
def interview_docx(tmp_path):
    """A small reviewed transcript: italic runs are Kiswahili."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("NRCCW KSM09 interview transcript")          # header, not a turn
    p = doc.add_paragraph("I: ")
    p.add_run("Habari yako").italic = True
    p.add_run(", how are you today?")
    p = doc.add_paragraph("R: ")
    p.add_run("Niko").italic = True
    p.add_run(" fine [unclear] ")
    p.add_run("nime").italic = True
    p.add_run("download the form in 2019.")
    doc.add_paragraph("and then I left")                           # continuation of R
    path = tmp_path / "interview.docx"
    doc.save(path)
    return path
