"""Build a ~4 minute mock interview from public FLEURS clips (known answers).

Use it as a smoke test before the real run: it downloads the FLEURS Kenyan
Kiswahili and US English dev sets (~310 MB, once), stitches 16 speaker turns
(some mixing Kiswahili and English), writes a matching DOCX with Kiswahili in
italics, and includes one turn that is written but never spoken. A healthy
pilot flags exactly that turn.

    python -m src.transcribe.pilot.make_mock --out data/pilot/mock
    python -m src.transcribe.pilot.run --audio data/pilot/mock/mock_interview.wav `
        --transcript data/pilot/mock/mock_interview.docx --out data/pilot/mock/run --languages sw
"""

import argparse
import io
import random
import tarfile
import urllib.request
from pathlib import Path

import numpy as np

SR = 16000
BASE = "https://huggingface.co/datasets/google/fleurs/resolve/main/data"
UNSPOKEN_TURN = 7


def fetch(url: str, path: Path) -> Path:
    if not path.exists():
        print(f"Downloading {url}")
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(path.suffix + ".part")
        urllib.request.urlretrieve(url, tmp)
        tmp.rename(path)
    return path


def pick(tsv: Path, n: int, rng: random.Random):
    seen, rows = set(), []
    for line in tsv.read_text(encoding="utf-8").splitlines():
        parts = line.split("\t")
        if parts[0] not in seen:
            seen.add(parts[0])
            rows.append((parts[1], parts[2]))
    rng.shuffle(rows)
    return rows[:n]


def load(archive: Path, items):
    import soundfile as sf

    wanted = {name for name, _ in items}
    clips = {}
    with tarfile.open(archive) as tar:
        for member in tar:
            name = member.name.rsplit("/", 1)[-1]
            if name in wanted:
                audio, sr = sf.read(io.BytesIO(tar.extractfile(member).read()), dtype="float32")
                if sr != SR:
                    raise ValueError(f"{name}: expected {SR} Hz, got {sr}")
                clips[name] = audio
                if len(clips) == len(wanted):
                    break
    return clips


def main() -> None:
    import soundfile as sf
    from docx import Document

    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default="data/pilot/mock")
    parser.add_argument("--seed", type=int, default=3)
    args = parser.parse_args()
    out = Path(args.out)
    cache = out / "fleurs"
    rng = random.Random(args.seed)

    items, clips = {}, {}
    for lang in ("sw_ke", "en_us"):
        tsv = fetch(f"{BASE}/{lang}/dev.tsv", cache / f"{lang}.dev.tsv")
        archive = fetch(f"{BASE}/{lang}/audio/dev.tar.gz", cache / f"{lang}.dev.tar.gz")
        items[lang] = pick(tsv, 14, rng)
        clips[lang] = load(archive, items[lang])
    sw, en = iter(items["sw_ke"]), iter(items["en_us"])

    audio = [np.zeros(SR, dtype=np.float32)]
    doc = Document()
    doc.add_paragraph("MOCK interview transcript (FLEURS dev clips)")
    for n in range(16):
        speaker = "I" if n % 2 == 0 else "R"
        if n == UNSPOKEN_TURN:
            parts = [("sw", next(sw)[1], None)]
        elif speaker == "I":
            name, text = next(en)
            parts = [("en", text, clips["en_us"][name])]
        elif n % 4 == 1:
            (s_name, s_text), (e_name, e_text) = next(sw), next(en)
            parts = [("sw", s_text, clips["sw_ke"][s_name]), ("en", e_text, clips["en_us"][e_name])]
        else:
            name, text = next(sw)
            parts = [("sw", text, clips["sw_ke"][name])]
        if n != UNSPOKEN_TURN:
            for _, _, clip in parts:
                audio += [clip, np.zeros(int(SR * 0.4), dtype=np.float32)]
            pause = 6.0 if n == 10 else rng.uniform(0.8, 1.8)
            audio.append(np.zeros(int(SR * pause), dtype=np.float32))
        paragraph = doc.add_paragraph(f"{speaker}: ")
        for i, (lang, text, _) in enumerate(parts):
            paragraph.add_run(("" if i == 0 else " ") + text).italic = lang == "sw"

    wav = out / "mock_interview.wav"
    sf.write(wav, np.concatenate(audio), SR)
    doc.save(out / "mock_interview.docx")
    print(f"Wrote {wav} and {out / 'mock_interview.docx'}")
    print(f"Expected: turn p{UNSPOKEN_TURN + 1:04d} (never spoken) is the only suspect turn.")


if __name__ == "__main__":
    main()
